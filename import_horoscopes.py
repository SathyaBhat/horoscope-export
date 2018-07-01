#!/usr/bin/env python

from __future__ import unicode_literals
from openpyxl import load_workbook
from sys import exit
import pymysql

import argparse

def create_tables(connection, sheet):
    pass

def create_db_connection(args):
    connection = pymysql.connect(host=args.host,
                                 user=args.user,
                                 passwd=args.password)
    return connection

def create_database(connection):
    try:
        connection.cursor().execute('create database horoscope')
    except pymysql.err.ProgrammingError as e:
        code, msg = e.args
        if code == 1007:
            print "Database exists, not recreating."
        else:
            print "{} {}".format(code, msg)
            exit(-1)

def export_to_mysql(file_name, connection):
    book = load_workbook(file_name)
    sheets = book.sheetnames
    values = []
    
    for sheet in sheets:
        create_tables(connection, sheet)

    for sheet in sheets:
        current_sheet = book.get_sheet_by_name(sheet)
        cells = current_sheet['A1': 'G32']
        headers = []
        
        for header in cells[0]:
            headers.append(header.value)
        for data in cells[1:]:
            values.append(data[0].value  + ' ' +  data[1].value.strftime('%Y-%m-%d') + ' ' + data[2].value + ' ' + data[3].value + ' ' + data[4].value + ' ' + data[5].value + ' ' + data[6].value)


if __name__=='__main__':
    parser = argparse.ArgumentParser(description='Script to parse the horoscope Excel.')
    parser.add_argument('--file', '-f', required=True, help='Filename to be read')
    parser.add_argument('--host', '-o', required=True, help='MySQL Host')
    parser.add_argument('--user', '-u', required=True, help='MySQL Username')
    parser.add_argument('--password', '-p', required=True, help='MySQL password')
    args = parser.parse_args()
    connection = create_db_connection(args)
    create_database(connection)
    export_to_mysql(args.file, connection)
